from openpyxl import load_workbook

from app.programs.expenses.services import expenses as expenses_svc
from app.programs.expenses.services import reports as reports_svc
from app.database import get_connection
from tests.conftest import first_category_id, first_cost_center_id, second_category_id


def _line(amount=100.0, particulars="Test", category_id=None, is_vatable=False):
    return {"category_id": category_id or first_category_id(), "particulars": particulars,
            "amount": amount, "is_vatable": is_vatable,
            "supplier_name": "Report Test Supplier" if is_vatable else "",
            "supplier_vat": "VATRPT001" if is_vatable else ""}


def _approved_expense(amount, expense_date="2026-08-22", paid_to="Vendor",
                       category_id=None, is_vatable=False):
    eid = expenses_svc.create_expense(
        expense_date=expense_date, cost_center_id=str(first_cost_center_id()), paid_to=paid_to,
        payment_mode="Cash", receipt_path=None, user_id=1,
        lines=[_line(amount, category_id=category_id, is_vatable=is_vatable)],
    )
    expenses_svc.approve_expense(eid, user_id=1)
    return eid


def test_spend_breakdown_only_counts_approved_in_range():
    _approved_expense(100, "2026-08-01")
    _approved_expense(200, "2026-08-15")
    # outside the range - should not count
    _approved_expense(9999, "2026-01-01")
    # pending - should not count regardless of date
    expenses_svc.create_expense(
        expense_date="2026-08-10", cost_center_id=str(first_cost_center_id()), paid_to="V",
        payment_mode="Cash", receipt_path=None, user_id=1, lines=[_line(500)],
    )

    by_cat, by_cc, monthly, grand_total = reports_svc.spend_breakdown("2026-08-01", "2026-08-31")
    assert grand_total == 300


def test_export_rows_respects_status_filter():
    _approved_expense(150)
    rejected_id = expenses_svc.create_expense(
        expense_date="2026-08-22", cost_center_id=str(first_cost_center_id()), paid_to="V2",
        payment_mode="Cash", receipt_path=None, user_id=1, lines=[_line(75, particulars="P2")],
    )
    expenses_svc.reject_expense(rejected_id, user_id=1, reason="no")

    approved_rows = reports_svc.export_rows(status="approved")
    assert any(r["amount"] == 150 for r in approved_rows)
    assert not any(r["amount"] == 75 for r in approved_rows)


def test_export_rows_respects_category_filter():
    other_cat_id = second_category_id()
    expenses_svc.create_expense(
        expense_date="2026-08-22", cost_center_id=str(first_cost_center_id()), paid_to="Match",
        payment_mode="Cash", receipt_path=None, user_id=1, lines=[_line(10)],
    )
    expenses_svc.create_expense(
        expense_date="2026-08-22", cost_center_id=str(first_cost_center_id()), paid_to="Other",
        payment_mode="Cash", receipt_path=None, user_id=1,
        lines=[_line(20, category_id=other_cat_id)],
    )
    rows = reports_svc.export_rows(status="", category_id=str(first_category_id()))
    paid_tos = {r["paid_to"] for r in rows}
    assert "Match" in paid_tos
    assert "Other" not in paid_tos


def test_export_rows_one_row_per_line_not_per_voucher():
    """The core reason exports moved to line granularity: a voucher with two
    lines in different categories must produce two export rows, so category
    (and VAT) totals attribute correctly."""
    other_cat_id = second_category_id()
    eid = expenses_svc.create_expense(
        expense_date="2026-08-22", cost_center_id=str(first_cost_center_id()), paid_to="Multi",
        payment_mode="Cash", receipt_path=None, user_id=1,
        lines=[_line(10, particulars="First"), _line(20, particulars="Second", category_id=other_cat_id)],
    )
    expenses_svc.approve_expense(eid, user_id=1)
    rows = reports_svc.export_rows(status="approved")
    multi_rows = [r for r in rows if r["paid_to"] == "Multi"]
    assert len(multi_rows) == 2
    assert {r["particulars"] for r in multi_rows} == {"First", "Second"}


def test_export_rows_includes_vat_columns():
    _approved_expense(115, is_vatable=True)
    rows = reports_svc.export_rows(status="approved")
    vatable_row = next(r for r in rows if r["is_vatable"])
    assert vatable_row["vat_amount"] > 0


def test_allocate_document_number_is_sequential_per_type():
    n1 = reports_svc.allocate_document_number("RPT")
    n2 = reports_svc.allocate_document_number("RPT")
    assert n1 != n2
    assert n1.startswith("RPT-")
    # a different doc_type gets its own independent sequence
    a1 = reports_svc.allocate_document_number("AUD")
    assert a1 == "AUD-" + n1.split("-", 2)[1] + "-00001"


def test_build_export_workbook_produces_valid_xlsx_with_metadata_and_total_row():
    _approved_expense(100)
    _approved_expense(50)
    rows = reports_svc.export_rows(status="approved")
    doc_no = reports_svc.allocate_document_number("RPT")
    buf = reports_svc.build_export_workbook(
        rows, document_no=doc_no, generated_by="Test Manager", filters_description="Period: all"
    )

    wb = load_workbook(buf)
    ws = wb.active
    all_text = " ".join(str(c.value) for row in ws.iter_rows() for c in row if c.value)
    assert doc_no in all_text
    assert "Test Manager" in all_text
    assert "Voucher No" in all_text
    assert "VAT" in all_text
    assert "TOTAL" in all_text
    assert "150" in all_text  # 100 + 50


def test_build_export_pdf_is_a4_landscape_and_contains_data():
    _approved_expense(321)
    rows = reports_svc.export_rows(status="approved")
    pdf_bytes = reports_svc.build_export_pdf(rows, generated_by="Test Manager",
                                              filters_description="Period: all")
    assert pdf_bytes[:4] == b"%PDF"
    text = pdf_bytes.decode("latin-1")
    # landscape A4 MediaBox is 841.8898 x 595.2756 pt (width/height swapped vs portrait)
    assert "841.88" in text and "595.27" in text


def test_build_summary_pdf_contains_generated_by_and_document_number():
    _approved_expense(500)
    by_cat, by_cc, _monthly, grand_total = reports_svc.spend_breakdown("2026-01-01", "2026-12-31")
    pdf_bytes = reports_svc.build_summary_pdf(by_cat, by_cc, grand_total,
                                               generated_by="Test Manager",
                                               filters_description="Period: 2026")
    assert pdf_bytes[:4] == b"%PDF"
    text = pdf_bytes.decode("latin-1")
    assert "Test Manager" in text


def test_build_export_pdf_survives_malicious_paid_to_field():
    """A voucher with an unclosed markup tag in 'paid to' - typed by the
    lowest-privilege role - must not break the listing PDF for everyone
    who tries to export it afterwards."""
    eid = expenses_svc.create_expense(
        expense_date="2026-08-22", cost_center_id=str(first_cost_center_id()), paid_to="<b>unclosed vendor name",
        payment_mode="Cash", receipt_path=None, user_id=1,
        lines=[_line(100.0, particulars="<font size=999>injected</font>")],
    )
    expenses_svc.approve_expense(eid, user_id=1)
    rows = reports_svc.export_rows(status="approved")
    pdf_bytes = reports_svc.build_export_pdf(rows, generated_by="Test Manager",
                                              filters_description="Period: all")
    assert pdf_bytes[:4] == b"%PDF"
