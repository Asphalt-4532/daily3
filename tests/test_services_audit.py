from app.services import audit as svc
from app.programs.expenses.services import settings as expense_settings_svc


def test_audit_log_records_actions():
    expense_settings_svc.add_category("Audited Category", acting_user_id=1)
    entries = svc.list_audit_log()
    assert any(e["action"] == "add_category" and e["details"] == "Audited Category" for e in entries)


def test_audit_log_date_filter_excludes_out_of_range_entries():
    from app.database import get_db
    expense_settings_svc.add_category("In Range Category", acting_user_id=1)
    with get_db() as conn:
        conn.execute("UPDATE audit_log SET created_at = '2020-01-01 00:00:00' "
                     "WHERE action='add_category' AND details='In Range Category'")
    expense_settings_svc.add_category("Also Recent", acting_user_id=1)

    recent_only = svc.list_audit_log(start="2026-01-01")
    details = [e["details"] for e in recent_only]
    assert "Also Recent" in details
    assert "In Range Category" not in details


def test_describe_filters_with_no_range():
    assert svc.describe_filters() == "All records"


def test_describe_filters_with_range():
    assert svc.describe_filters("2026-01-01", "2026-12-31") == "Period: 2026-01-01 to 2026-12-31"


def test_allocate_document_number_is_sequential():
    n1 = svc.allocate_document_number("AUD")
    n2 = svc.allocate_document_number("AUD")
    assert n1 != n2
    assert n1.startswith("AUD-")


def test_build_audit_export_workbook_contains_metadata_and_entries():
    from openpyxl import load_workbook
    expense_settings_svc.add_category("Exported Category", acting_user_id=1)
    entries = svc.list_audit_log(limit=0)
    buf = svc.build_audit_export_workbook(
        entries, document_no="AUD-2026-00001", generated_by="Test Manager",
        filters_description="All records",
    )
    wb = load_workbook(buf)
    ws = wb.active
    all_text = " ".join(str(c.value) for row in ws.iter_rows() for c in row if c.value)
    assert "AUD-2026-00001" in all_text
    assert "Test Manager" in all_text
    assert "add_category" in all_text


def test_build_audit_export_pdf_is_valid_and_landscape():
    expense_settings_svc.add_category("PDF Export Category", acting_user_id=1)
    entries = svc.list_audit_log(limit=0)
    pdf_bytes = svc.build_audit_export_pdf(entries, generated_by="Test Manager",
                                            filters_description="All records")
    assert pdf_bytes[:4] == b"%PDF"
    text = pdf_bytes.decode("latin-1")
    assert "841.88" in text and "595.27" in text  # landscape A4 MediaBox


def test_build_audit_export_pdf_survives_malformed_markup_in_details():
    """A rejection/delete reason is free text that lands in audit_log.details
    (often inside a JSON snapshot) and gets rendered into the audit PDF.
    An unclosed tag there must not crash the export - proving the shared
    pdf_common escaping still applies after the audit builder moved out of
    the Expense Program's pdf_reports.py."""
    from app.database import get_db
    with get_db() as conn:
        conn.execute(
            "INSERT INTO audit_log (user_id, action, details) VALUES (1, 'reject_expense', ?)",
            ('{"voucher_no": "PCV-2026-00099", "reason": "<b>unclosed reason tag"}',),
        )
    entries = svc.list_audit_log(limit=0)
    pdf_bytes = svc.build_audit_export_pdf(entries, generated_by="Test Manager",
                                            filters_description="All records")
    assert pdf_bytes[:4] == b"%PDF"
